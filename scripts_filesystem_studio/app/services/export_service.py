"""
services/export_service.py — Exportación unificada de resultados.

Sustituye a los serializadores artesanales dispersos en los scripts
originales (CSV/JSON del detector, Markdown del reporte, JSON del
creator). Un solo lugar para: CSV, JSON, Markdown, HTML, PDF y Excel.

PDF se genera con QTextDocument + QPdfWriter (sin dependencias extra).
Excel usa openpyxl si está instalado; si no, se lanza un error claro.
"""

import json
from pathlib import Path

from PySide6.QtCore import QMarginsF
from PySide6.QtGui import QPageLayout, QPageSize, QPdfWriter, QTextDocument


def _ensure_parent(path: str) -> None:
    Path(path).parent.mkdir(parents=True, exist_ok=True)


def save_text(text: str, path: str) -> None:
    _ensure_parent(path)
    Path(path).write_text(text, encoding="utf-8")


def save_json(data, path: str) -> None:
    _ensure_parent(path)
    with open(path, "w", encoding="utf-8") as f:
        json.dump(data, f, indent=2, ensure_ascii=False)


def save_csv(headers: list[str], rows: list[list], path: str) -> None:
    import csv
    _ensure_parent(path)
    with open(path, "w", newline="", encoding="utf-8") as f:
        writer = csv.writer(f)
        writer.writerow(headers)
        writer.writerows(rows)


def save_markdown_table(title: str, headers: list[str], rows: list[list],
                        path: str, intro: str = "") -> None:
    lines = [f"# {title}", ""]
    if intro:
        lines += [intro, ""]
    lines.append("| " + " | ".join(headers) + " |")
    lines.append("| " + " | ".join("---" for _ in headers) + " |")
    for row in rows:
        cells = [str(cell).replace("|", "\\|") for cell in row]
        lines.append("| " + " | ".join(cells) + " |")
    save_text("\n".join(lines) + "\n", path)


def save_html(html: str, path: str) -> None:
    save_text(html, path)


def markdown_to_html(markdown_text: str, title: str = "") -> str:
    """Convierte Markdown a HTML usando el soporte nativo de Qt."""
    doc = QTextDocument()
    doc.setMarkdown(markdown_text)
    body = doc.toHtml()
    if title:
        body = body.replace("<title></title>", f"<title>{title}</title>")
    return body


def save_pdf_from_html(html: str, path: str) -> None:
    """Renderiza HTML a PDF con el motor de texto enriquecido de Qt."""
    _ensure_parent(path)
    writer = QPdfWriter(path)
    writer.setPageSize(QPageSize(QPageSize.A4))
    writer.setPageMargins(QMarginsF(15, 15, 15, 15), QPageLayout.Millimeter)
    doc = QTextDocument()
    doc.setHtml(html)
    doc.setPageSize(writer.pageLayout().paintRectPixels(
        writer.resolution()).size().toSizeF())
    doc.print_(writer)


def save_excel(headers: list[str], rows: list[list], path: str,
               sheet_title: str = "Datos") -> None:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font
    except ImportError as exc:
        raise RuntimeError(
            "Exportar a Excel requiere el paquete 'openpyxl' "
            "(pip install openpyxl).") from exc

    _ensure_parent(path)
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_title
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for row in rows:
        ws.append(row)
    # Ancho de columnas razonable según contenido
    for col_idx, header in enumerate(headers, start=1):
        width = max([len(str(header))] +
                    [len(str(r[col_idx - 1])) for r in rows[:200]]) + 2
        ws.column_dimensions[ws.cell(row=1, column=col_idx)
                             .column_letter].width = min(width, 60)
    wb.save(path)
