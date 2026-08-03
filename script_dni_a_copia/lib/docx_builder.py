"""Word document assembly and PDF export for dni-a-copia.

Builds an A4 page with both faces centered horizontally and vertically at
real ID-1 size, then (optionally) exports it to PDF via LibreOffice
headless. Images are embedded from memory (BytesIO), so no intermediate
files are needed for the Word step.
"""
from __future__ import annotations

import io
import shutil
import subprocess
import sys
import tempfile
from logging import Logger
from pathlib import Path
from typing import Optional

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt
from PIL import Image

import config


def _enable_vertical_center(section) -> None:
    """Sets w:vAlign=center on the section (honored by MS Word).

    Must be inserted in schema order (before w:docGrid), or processors
    silently ignore it. LibreOffice ignores it regardless, so the symmetric
    top/bottom margins below are what actually center the content there.
    """
    sect_pr = section._sectPr
    for existing in sect_pr.findall(qn("w:vAlign")):
        sect_pr.remove(existing)
    valign = OxmlElement("w:vAlign")
    valign.set(qn("w:val"), "center")
    reference = None
    for tag in ("w:docGrid", "w:textDirection", "w:bidi", "w:rtlGutter"):
        reference = sect_pr.find(qn(tag))
        if reference is not None:
            break
    reference.addprevious(valign) if reference is not None \
        else sect_pr.append(valign)


def _configure_page(section, geom: config.Geometry) -> None:
    """A4 portrait with margins that center both faces vertically."""
    section.page_width = Mm(config.PAGE_W_MM)
    section.page_height = Mm(config.PAGE_H_MM)
    content_h = 2 * geom.img_h_mm + config.GAP_MM
    section.top_margin = section.bottom_margin = \
        Mm((config.PAGE_H_MM - content_h) / 2)
    section.left_margin = section.right_margin = Mm(config.SIDE_MARGIN_MM)
    _enable_vertical_center(section)


def _add_face(doc: Document, image: Image.Image,
              geom: config.Geometry) -> None:
    """Adds one centered, real-size face image (embedded from memory)."""
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(0)
    fmt.space_after = Pt(0)
    fmt.line_spacing_rule = WD_LINE_SPACING.SINGLE
    buffer = io.BytesIO()
    image.save(buffer, format="PNG")
    buffer.seek(0)
    paragraph.add_run().add_picture(
        buffer, width=Mm(geom.img_w_mm), height=Mm(geom.img_h_mm))


def build_document(front: Image.Image, back: Image.Image,
                   geom: config.Geometry, out_path: Path,
                   dry_run: bool, logger: Logger) -> Path:
    """Builds the Word document and saves it (unless dry-run)."""
    if dry_run:
        logger.info("[SIMULACIÓN] Se generaría el Word '%s'.", out_path)
        return out_path
    doc = Document()
    _configure_page(doc.sections[0], geom)
    _add_face(doc, front, geom)
    doc.paragraphs[-1].paragraph_format.space_after = Mm(config.GAP_MM)
    _add_face(doc, back, geom)
    doc.save(str(out_path))
    logger.info("Word guardado: %s", out_path)
    return out_path


def export_pdf(docx_path: Path, out_dir: Path,
               dry_run: bool, logger: Logger) -> Optional[Path]:
    """Converts the .docx to PDF with LibreOffice headless.

    Uses a throwaway user profile so it never clashes with a running
    LibreOffice instance. Returns the PDF path, or None in dry-run.
    """
    pdf_path = out_dir / f"{docx_path.stem}.pdf"
    if dry_run:
        logger.info("[SIMULACIÓN] Se exportaría el PDF '%s'.", pdf_path)
        return None
    soffice = next((shutil.which(name) for name in config.SOFFICE_CANDIDATES
                    if shutil.which(name)), None)
    if soffice is None:
        logger.error("No se encontró LibreOffice para exportar a PDF.")
        sys.exit(config.EXIT_DEPENDENCY)
    profile = Path(tempfile.mkdtemp(prefix="dni_a_copia_lo_"))
    try:
        _run_soffice(soffice, profile, docx_path, out_dir, logger)
    finally:
        shutil.rmtree(profile, ignore_errors=True)
    logger.info("PDF guardado: %s", pdf_path)
    return pdf_path


def _run_soffice(soffice: str, profile: Path, docx_path: Path,
                 out_dir: Path, logger: Logger) -> None:
    """Runs the LibreOffice conversion, failing loudly on error."""
    command = [
        soffice, "--headless",
        f"-env:UserInstallation=file://{profile}",
        "--convert-to", config.PDF_EXPORT_FILTER,
        "--outdir", str(out_dir), str(docx_path),
    ]
    try:
        subprocess.run(command, check=True, capture_output=True, text=True)
    except subprocess.CalledProcessError as error:
        logger.error("LibreOffice falló al convertir a PDF: %s",
                     error.stderr.strip() or error)
        sys.exit(config.EXIT_GENERAL)
