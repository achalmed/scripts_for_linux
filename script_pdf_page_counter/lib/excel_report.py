"""
lib/excel_report.py — Generación del reporte Excel.

Construye el libro con una hoja de datos (bloques por blog con
subtotales) y una hoja de metadatos. Dividido en funciones pequeñas
para que cada bloque visual sea testeable por separado.

Author : Edison Achalma (@achalmed)
Version: 2.0.0
"""

import sys
from datetime import datetime
from typing import Dict, List

from config import EXIT_MISSING_DEPENDENCY
from lib.scanner import ESTADO_OK, ResultadoPdf

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.worksheet.worksheet import Worksheet
except ImportError:
    print("❌ Error: openpyxl no está instalado.")
    print("   Instala con: pip install openpyxl")
    print("   o bien:      conda install -c conda-forge openpyxl")
    sys.exit(EXIT_MISSING_DEPENDENCY)

# Paleta y estilos compartidos por todas las funciones de esta hoja
_HEADER_FONT = Font(bold=True, size=12, color="FFFFFF")
_HEADER_FILL = PatternFill(start_color="2E5090", end_color="2E5090", fill_type="solid")
_HEADER_ALIGNMENT = Alignment(horizontal="center", vertical="center")
_BLOG_FONT = Font(bold=True, size=11, color="FFFFFF")
_BLOG_FILL = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
_TOTAL_FONT = Font(bold=True, size=11)
_TOTAL_FILL = PatternFill(start_color="E7E6E6", end_color="E7E6E6", fill_type="solid")
_THIN_BORDER = Border(
    left=Side(style="thin"), right=Side(style="thin"),
    top=Side(style="thin"), bottom=Side(style="thin"),
)


def _write_column_headers(sheet: Worksheet) -> None:
    """Writes and styles the fixed column headers on row 1."""
    headers = {"A1": "Blog", "B1": "Ruta del Archivo",
               "C1": "Número de Páginas", "D1": "Estado"}
    for cell_ref, text in headers.items():
        sheet[cell_ref] = text
        sheet[cell_ref].font = _HEADER_FONT
        sheet[cell_ref].fill = _HEADER_FILL
        sheet[cell_ref].alignment = _HEADER_ALIGNMENT
        sheet[cell_ref].border = _THIN_BORDER


def _write_blog_block(sheet: Worksheet, row: int, blog_name: str,
                      results: List[ResultadoPdf]) -> (int, int):
    """
    Writes one blog's banner, data rows, and subtotal.

    Args:
        sheet: Target worksheet.
        row: First free row.
        blog_name: Display name of the blog.
        results: Scan results for this blog.

    Returns:
        Tuple (next_free_row, blog_page_subtotal).
    """
    sheet[f"A{row}"] = blog_name.upper()
    sheet[f"A{row}"].font = _BLOG_FONT
    sheet[f"A{row}"].fill = _BLOG_FILL
    sheet[f"A{row}"].alignment = Alignment(horizontal="left", vertical="center")
    sheet.merge_cells(f"A{row}:D{row}")
    row += 1

    subtotal_pages = 0
    for relative_path, page_count, status in results:
        sheet[f"B{row}"] = relative_path
        sheet[f"C{row}"] = page_count if status == ESTADO_OK else 0
        sheet[f"C{row}"].alignment = Alignment(horizontal="center")
        sheet[f"D{row}"] = status
        sheet[f"D{row}"].alignment = Alignment(horizontal="center")
        if status == ESTADO_OK:
            subtotal_pages += page_count
        row += 1

    sheet[f"B{row}"] = f"SUBTOTAL {blog_name}"
    sheet[f"B{row}"].font = Font(bold=True, italic=True)
    sheet[f"C{row}"] = subtotal_pages
    sheet[f"C{row}"].font = Font(bold=True, italic=True)
    sheet[f"C{row}"].alignment = Alignment(horizontal="center")
    sheet[f"C{row}"].fill = _TOTAL_FILL
    sheet[f"D{row}"] = f"{len(results)} archivos"
    sheet[f"D{row}"].font = Font(italic=True)
    sheet[f"D{row}"].alignment = Alignment(horizontal="center")

    return row + 2, subtotal_pages  # fila en blanco de separación entre blogs


def _write_grand_total(sheet: Worksheet, row: int,
                       total_pages: int, total_files: int) -> None:
    """Writes the final TOTAL GENERAL row."""
    sheet[f"B{row}"] = "TOTAL GENERAL"
    sheet[f"B{row}"].font = _TOTAL_FONT
    sheet[f"C{row}"] = total_pages
    sheet[f"C{row}"].font = _TOTAL_FONT
    sheet[f"C{row}"].alignment = Alignment(horizontal="center")
    sheet[f"C{row}"].fill = _TOTAL_FILL
    sheet[f"D{row}"] = f"{total_files} archivos"
    sheet[f"D{row}"].font = _TOTAL_FONT
    sheet[f"D{row}"].alignment = Alignment(horizontal="center")


def _write_metadata_sheet(workbook: Workbook, solo_index: bool,
                          blog_count: int, total_files: int,
                          total_pages: int) -> None:
    """Adds the 'Información' sheet with report metadata."""
    info = workbook.create_sheet("Información")
    info["A1"] = "Información del Reporte"
    info["A1"].font = Font(bold=True, size=14)
    rows = [
        ("Fecha de generación:", datetime.now().strftime("%d/%m/%Y %H:%M:%S")),
        ("Tipo de búsqueda:", "Solo index.pdf" if solo_index else "Todos los PDFs"),
        ("Total de blogs procesados:", blog_count),
        ("Total de archivos:", total_files),
        ("Total de páginas:", total_pages),
        ("Generado por:", "Edison Achalma - PDF Page Counter"),
    ]
    for offset, (label, value) in enumerate(rows, start=3):
        info[f"A{offset}"] = label
        info[f"B{offset}"] = value


def create_excel_report(results_by_blog: Dict[str, List[ResultadoPdf]],
                        output_path: str, solo_index: bool) -> None:
    """
    Builds and saves the complete Excel report.

    Args:
        results_by_blog: {blog_display_name: scan results}.
        output_path: Destination .xlsx path.
        solo_index: Whether only index.pdf files were scanned.
    """
    workbook = Workbook()
    data_sheet = workbook.active
    data_sheet.title = "Conteo de Páginas"

    _write_column_headers(data_sheet)

    current_row = 2
    total_pages = 0
    total_files = 0
    for blog_name, results in sorted(results_by_blog.items()):
        if not results:
            continue
        current_row, blog_subtotal = _write_blog_block(
            data_sheet, current_row, blog_name, results)
        total_pages += blog_subtotal
        total_files += len(results)

    _write_grand_total(data_sheet, current_row, total_pages, total_files)

    for column, width in (("A", 25), ("B", 70), ("C", 18), ("D", 15)):
        data_sheet.column_dimensions[column].width = width

    _write_metadata_sheet(workbook, solo_index, len(results_by_blog),
                          total_files, total_pages)
    workbook.save(output_path)
